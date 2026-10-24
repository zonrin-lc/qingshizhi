# -*- coding: utf-8 -*-
"""生成软著软件说明书 docx（图文）。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = r"D:\MUSI\QINGSHIZHI"
SHOTS = os.path.join(ROOT, "shots")
OUT = os.path.join(ROOT, "docs", "软著-软件说明书.docx")
SW = "青史志互动叙事游戏软件"
VER = "V1.0"

doc = Document()
sec = doc.sections[0]
sec.page_height = Pt(1169); sec.page_width = Pt(826)
sec.top_margin = Pt(72); sec.bottom_margin = Pt(72)
sec.left_margin = Pt(80); sec.right_margin = Pt(80)

# 页眉页脚
hdr = sec.header.paragraphs[0]
hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = hdr.add_run(f"{SW}  {VER}  软件说明书")
r.font.size = Pt(9)
ftr = sec.footer.paragraphs[0]
ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = ftr.add_run(); r.font.size = Pt(9)
from docx.oxml import OxmlElement
def add_field(paragraph, field):
    r = paragraph.add_run()
    a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'),'begin')
    b = OxmlElement('w:instrText'); b.set(qn('xml:space'),'preserve'); b.text = field
    c = OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'),'separate')
    d = OxmlElement('w:t'); d.text = "1"
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),'end')
    r._r.append(a); r._r.append(b); r._r.append(c); r._r.append(d); r._r.append(e)
add_field(ftr, "PAGE")

def set_font(run, size=12, bold=False, hei=False, color=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size); run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), "黑体" if hei else "宋体")
    if color: run.font.color.rgb = color

def h1(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(8)
    set_font(p.add_run(text), 16, True, hei=True)
    return p

def h2(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run(text), 13, True, hei=True)
    return p

def body(text):
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5; p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(text), 12)
    return p

def pic(name, caption, width_cm=12.5):
    path = os.path.join(SHOTS, name)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(width_cm))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)
    set_font(c.add_run(caption), 10.5, False, color=RGBColor(0x55,0x55,0x55))

# ============ 封面标题 ============
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.paragraph_format.space_before = Pt(60)
set_font(t.add_run(SW), 22, True, hei=True)
t2 = doc.add_paragraph(); t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(t2.add_run(f"软件说明书（{VER}）"), 16, True, hei=True)
t3 = doc.add_paragraph(); t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
t3.paragraph_format.space_after = Pt(24)
set_font(t3.add_run("文档版本：V1.0    编写日期：2026年10月21日"), 11)

# ============ 一、软件概述 ============
h1("一、软件概述")
body(f"{SW}（以下简称“本软件”）是一款以中国历史为题材的单机互动叙事游戏软件，以 HTML5 技术实现，可在移动端与桌面端浏览器及 TapTap 客户端中直接运行。")
body("本软件以“执笔入史”为核心概念：玩家在秦、汉、唐、宋、明五个历史朝代的剧本中扮演李斯、荆轲、韩信、项羽、陈胜等历史人物，从其人生的关键处境切入，通过阅读画卷式剧情、在关键节点做出抉择，影响体魄、财力、武力、文才、谋略五维属性与财富、名声、权势、政绩、君心、危机等身外属性的消长，最终走向不同的历史结局。")
body("本软件包含五个可玩剧本（仓鼠之局、图穷之局、兵仙之局、霸王之局、首义之局），共数十种结局，并提供青史图鉴收集系统、三档难度（剧情、普通、硬核）、剧情朗读、背景音乐与音效、本地存档与 TapTap 云存档同步等功能。软件整体采用国风水墨美术风格，界面以画卷、印章、罗盘等传统元素构成。")

# ============ 二、运行环境 ============
h1("二、运行环境")
h2("2.1 硬件环境")
body("客户端：智能手机、平板电脑或个人计算机。处理器主频 1.5GHz 及以上，运行内存 2GB 及以上，支持 WebGL 图形加速，屏幕分辨率不低于 960×540。")
body("服务端：本软件为单机游戏，除 TapTap 云存档同步功能依赖 TapTap 提供的云端服务外，无需自建服务器。")
h2("2.2 软件环境")
body("客户端操作系统：Android 8.0 及以上、iOS 13 及以上、Windows 10 及以上或 macOS。")
body("运行载体：Safari、Chrome、Edge 等现代浏览器，或 TapTap 客户端内嵌 WebView。要求浏览器支持 HTML5、WebGL、Web Audio 与 WebM 视频解码。")
body("开发语言与工具：TypeScript、JavaScript、HTML5、CSS3；构建工具 Vite；游戏规则引擎为自研 JavaScript 引擎。")

# ============ 三、安装与启动 ============
h1("三、安装与启动")
body("本软件为 H5 游戏，通过 TapTap 平台以 H5 包形式发布，玩家在 TapTap 客户端内点击游戏图标即可直接启动，无需传统意义的安装过程。")
body("启动后软件自动加载字体、音频与视频资源并进入主界面。首次加载后资源会由浏览器缓存，再次进入时可离线运行；云存档同步功能需在联网状态下使用。")

# ============ 四、功能说明 ============
h1("四、功能说明")

h2("4.1 剧本选择")
body("启动软件后进入剧本选择页。页面左侧为朝代罗盘，玩家可旋转选择秦、汉、唐、宋、明五个朝代；选中朝代后，右侧展示该朝代的已开卷局面立绘与局面卡片。每张立绘对应一个具体局面（如李斯的“仓鼠之局”、荆轲的“图穷之局”等）。页面右侧提供难度选择：剧情、普通、硬核三档，硬核难度下属性变化更严苛且每章可回溯次数受限。确认后点击“开始”进入对局。")
pic("select-mobile-drawer.png", "图 1  剧本选择页：朝代罗盘、局面立绘与难度选择")

h2("4.2 剧本预告")
body("对于尚未开卷的朝代，选择后展示该朝代的剧本预告页，包括朝代年代跨度、时代简介、本卷关键词标签以及即将登场的局面名称，供玩家了解后续内容。")
pic("select-mobile-teaser.png", "图 2  未开卷朝代的剧本预告页")

h2("4.3 对局主界面")
body("对局主界面由三部分组成：顶部为局面名称、当前章节与“回溯”按钮；中部为画卷式剧情正文区，左侧展示当前角色立绘；右侧为属性区，上方为体魄、财力、武力、文才、谋略五维属性的雷达图，下方为财富、名声、权势、政绩、君心、危机等身外属性进度条，以及偏离度与角色年龄。")
body("剧情正文以卷轴形式呈现，玩家阅读完毕后点击“继续”按钮推进剧情。当剧情出现抉择点时，正文下方出现事件或抉择选项卡，玩家点选后剧情分支推进；“回溯”按钮可消耗回溯次数返回本章之前的节点重新选择。")
pic("game-mobile-scroll-border.png", "图 3  对局主界面：画卷剧情正文与右侧属性区")
pic("game-intro-settled.png", "图 4  剧情落定后的正文与继续按钮")

h2("4.4 属性变化提示")
body("玩家的抉择会引起属性增减。属性发生变化时，软件先以浮层提示逐条展示变化内容（如“【探针】测试数据值增减——体魄+1”），提示队列播放完毕后，变化信息在正文中原位淡入留存，使玩家既能即时感知变化，又不打断阅读节奏。")
pic("game-intro-toast.png", "图 5  属性增减的浮层提示")

h2("4.5 结局结算")
body("剧情推进至终章后，软件根据玩家全程抉择与属性状态判定结局，弹出结局结算卡。结算卡展示结局印鉴（如“败局”）、结局名称与天干编号、“太史公曰”史评体结语，以及该结局的累计达成次数。玩家点击“返回主页”可回到剧本选择页，重新体验其他分支。")
pic("ending-mobile.png", "图 6  结局结算页")

h2("4.6 青史图鉴")
body("青史图鉴收录玩家达成过的全部结局。列表页左侧为人物立绘与生平跨度，右侧上方为收集总览进度（如 1/82），并按局面分类展示印章收集进度；下方以卡片列出各结局，已达成的结局显示印鉴与达成次数，未达成的结局以“未达成”占位。点击已达成结局可查看详情，包括结局印鉴、太史公曰完整评语与达成次数。")
pic("gallery-list-mobile.png", "图 7  青史图鉴列表页")
pic("galdetail-mobile.png", "图 8  图鉴结局详情页")

h2("4.7 设置")
body("设置页提供背景音乐、游戏音效、剧情朗读三项音量调节；提供云存档同步功能（在 TapTap 客户端内可用，可将本地存档上传云端或从云端下载）；页面底部显示当前软件版本号。")
pic("settings-version.png", "图 9  设置页")

# ============ 五、操作说明 ============
h1("五、操作说明")
body("本软件全部操作基于点击或触摸完成：点按按钮、选项卡与滑杆进行交互；罗盘通过点选朝代图标切换；音量滑杆通过拖动调节；画卷正文区域支持滚动阅读长剧情；各页面的返回按钮位于左上角。")

# ============ 六、注意事项 ============
h1("六、注意事项")
body("云存档同步功能仅在 TapTap 客户端环境中可用，浏览器直接运行时为本地存档。剧情朗读功能依赖设备音频输出，请在体验时开启声音。硬核难度下回溯次数受限，请谨慎抉择。")

doc.save(OUT)
print("saved:", OUT)
